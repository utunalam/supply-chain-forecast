
import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from database import SessionLocal
from models import InventoryRecord, SalesRecord


PROJECT_ROOT = Path(__file__).resolve().parent.parent
INCOMING_DIR = PROJECT_ROOT / "data" / "incoming"

INVENTORY_FILE = INCOMING_DIR / "Inventory_Report_SYNTHETIC.csv"
SALES_FILE = INCOMING_DIR / "Sales_QTY_Breakdown_SYNTHETIC.xlsx"


def parse_datetime(value):
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value

    raise ValueError(f"Unsupported date value: {value!r}")


def parse_snapshot_date(value):
    if value is None or str(value).strip() == "":
        raise ValueError("Snapshot_Date is blank")

    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise ValueError(
            f"Invalid Snapshot_Date: {value!r}. "
            "Expected YYYY-MM-DD."
        ) from exc


def parse_integer(value, field_name):
    if value is None or str(value).strip() == "":
        raise ValueError(f"{field_name} is blank")

    result = int(value)

    if result < 0:
        raise ValueError(f"{field_name} cannot be negative")

    return result


def parse_float(value, field_name):
    if value is None or str(value).strip() == "":
        return None

    result = float(value)

    if result < 0:
        raise ValueError(f"{field_name} cannot be negative")

    return result


def import_inventory(db):
    source_file = INVENTORY_FILE.name
    imported = 0
    updated = 0
    skipped = 0

    # Ignore blank lines before or between CSV rows.
    with INVENTORY_FILE.open(
        encoding="utf-8-sig", newline=""
    ) as file:
        reader = csv.DictReader(
            line for line in file if line.strip()
        )

        required = {
            "Item_Key",
            "Partner_Item_ID (12NC)",
            "item_description",
            "Partner",
            "AVAILABLE_QTY",
            "Expire_Date",
            "work_zone",
            "Snapshot_Date",
            "Synthetic_Data_Flag",
        }

        if not reader.fieldnames:
            raise ValueError("Inventory CSV has no header")

        missing = required - set(reader.fieldnames)

        if missing:
            raise ValueError(
                f"Inventory CSV missing columns: {sorted(missing)}"
            )

        for row_number, row in enumerate(reader, start=2):
            # Parse and validate the row before modifying the database.
            snapshot_date = parse_snapshot_date(
                row["Snapshot_Date"]
            )
            available_qty = parse_integer(
                row["AVAILABLE_QTY"], "AVAILABLE_QTY"
            )

            inventory_item_id = row[
                "Partner_Item_ID (12NC)"
            ].strip()

            if not inventory_item_id:
                raise ValueError(
                    f"Inventory item ID is blank at CSV row {row_number}"
                )

            item_key = row["Item_Key"].strip()
            item_description = row["item_description"].strip()
            partner = row["Partner"].strip() or None
            expire_date = row["Expire_Date"].strip() or None
            work_zone = row["work_zone"].strip() or None
            source_flag = (
                row["Synthetic_Data_Flag"].strip() or None
            )

            existing = (
                db.query(InventoryRecord)
                .filter_by(
                    source_file=source_file,
                    source_row_number=row_number,
                )
                .first()
            )

            values = {
                "item_key": item_key,
                "inventory_item_id": inventory_item_id,
                "item_description": item_description,
                "partner": partner,
                "available_qty": available_qty,
                "expire_date": expire_date,
                "work_zone": work_zone,
                "snapshot_date": snapshot_date,
                "source_flag": source_flag,
            }

            if existing:
                for field, value in values.items():
                    setattr(existing, field, value)

                updated += 1
            else:
                record = InventoryRecord(
                    **values,
                    source_file=source_file,
                    source_row_number=row_number,
                )
                db.add(record)
                imported += 1

    return imported, updated, skipped


def import_sales(db):
    source_file = SALES_FILE.name
    imported = 0
    skipped = 0

    workbook = load_workbook(
        SALES_FILE, read_only=True, data_only=True
    )

    try:
        if "Sales Data" not in workbook.sheetnames:
            raise ValueError("Worksheet 'Sales Data' not found")

        sheet = workbook["Sales Data"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)

        positions = {
            str(name).strip(): index
            for index, name in enumerate(header)
            if name is not None
        }

        required = {
            "Order Date",
            "Completed Date",
            "Status",
            "Warehouse",
            "Country",
            "Sales Channel",
            "Partner Order ID",
            "aCommerce Order ID",
            "Product SKU",
            "Item ID",
            "Item Name",
            "Item Quantity",
            "Selling Price",
            "Synthetic_Data_Flag",
        }

        missing = required - set(positions)

        if missing:
            raise ValueError(
                f"Sales workbook missing columns: {sorted(missing)}"
            )

        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None for value in row):
                continue

            existing = (
                db.query(SalesRecord)
                .filter_by(
                    source_file=source_file,
                    source_row_number=row_number,
                )
                .first()
            )

            # Existing sales rows are deliberately left unchanged.
            if existing:
                skipped += 1
                continue

            def value(column):
                index = positions[column]
                return row[index] if index < len(row) else None

            quantity = parse_integer(
                value("Item Quantity"), "Item Quantity"
            )

            record = SalesRecord(
                order_date=parse_datetime(value("Order Date")),
                completed_date=parse_datetime(
                    value("Completed Date")
                ),
                product_sku=(
                    str(value("Product SKU")).strip()
                    if value("Product SKU") is not None
                    else None
                ),
                sales_item_id=(
                    str(value("Item ID")).strip()
                    if value("Item ID") is not None
                    else None
                ),
                item_name=(
                    str(value("Item Name")).strip()
                    if value("Item Name") is not None
                    else None
                ),
                item_quantity=quantity,
                selling_price=parse_float(
                    value("Selling Price"), "Selling Price"
                ),
                order_status=(
                    str(value("Status")).strip()
                    if value("Status") is not None
                    else None
                ),
                warehouse=(
                    str(value("Warehouse")).strip()
                    if value("Warehouse") is not None
                    else None
                ),
                country=(
                    str(value("Country")).strip()
                    if value("Country") is not None
                    else None
                ),
                sales_channel=(
                    str(value("Sales Channel")).strip()
                    if value("Sales Channel") is not None
                    else None
                ),
                source_flag=(
                    str(value("Synthetic_Data_Flag")).strip()
                    if value("Synthetic_Data_Flag") is not None
                    else None
                ),
                source_file=source_file,
                source_row_number=row_number,
                partner_order_id=(
                    str(value("Partner Order ID")).strip()
                    if value("Partner Order ID") is not None
                    else None
                ),
                acommerce_order_id=(
                    str(value("aCommerce Order ID")).strip()
                    if value("aCommerce Order ID") is not None
                    else None
                ),
            )

            db.add(record)
            imported += 1

    finally:
        workbook.close()

    return imported, skipped


def main():
    if not INVENTORY_FILE.exists():
        raise FileNotFoundError(INVENTORY_FILE)

    if not SALES_FILE.exists():
        raise FileNotFoundError(SALES_FILE)

    db = SessionLocal()

    try:
        inventory_imported, inventory_updated, inventory_skipped = (
            import_inventory(db)
        )

        sales_imported, sales_skipped = import_sales(db)

        db.commit()

        print("Report import completed successfully.")
        print(
            f"Inventory: {inventory_imported} imported, "
            f"{inventory_updated} updated, "
            f"{inventory_skipped} skipped"
        )
        print(
            f"Sales: {sales_imported} imported, "
            f"{sales_skipped} skipped"
        )
        print("Snapshot dates were stored for inventory rows.")
        print("Existing SKUs and product mappings were not modified.")

    except Exception:
        db.rollback()
        raise

    finally:
        db.close()


if __name__ == "__main__":
    main()
