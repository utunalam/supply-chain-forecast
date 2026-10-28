
import csv
from pathlib import Path
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
INCOMING_DIR = PROJECT_ROOT / "data" / "incoming"

INVENTORY_FILE = INCOMING_DIR / "Inventory_Report_SYNTHETIC.csv"
SALES_FILE = INCOMING_DIR / "Sales_QTY_Breakdown_SYNTHETIC.xlsx"


def validate_inventory():
    required_columns = {
        "Item_Key",
        "Partner_Item_ID (12NC)",
        "item_description",
        "AVAILABLE_QTY",
        "Synthetic_Data_Flag",
    }

    with INVENTORY_FILE.open(
        encoding="utf-8-sig", newline=""
    ) as file:
        reader = csv.DictReader(file)

        if not reader.fieldnames:
            raise ValueError("Inventory CSV has no header.")

        missing = required_columns - set(reader.fieldnames)
        if missing:
            raise ValueError(
                f"Inventory CSV missing columns: {sorted(missing)}"
            )

        total = 0
        errors = []

        for row_number, row in enumerate(reader, start=2):
            try:
                item_key = row["Item_Key"].strip()
                item_id = row["Partner_Item_ID (12NC)"].strip()
                description = row["item_description"].strip()
                quantity = int(row["AVAILABLE_QTY"].strip())

                if not item_key or not item_id or not description:
                    raise ValueError("Required field is blank.")

                if quantity < 0:
                    raise ValueError("Available quantity is negative.")

                total += 1

            except (ValueError, AttributeError) as exc:
                errors.append(f"Row {row_number}: {exc}")

    print(f"Inventory rows checked: {total + len(errors)}")
    print(f"Inventory valid rows: {total}")
    print(f"Inventory errors: {len(errors)}")

    for error in errors:
        print(error)

    return total, len(errors)


def validate_sales():
    required_columns = {
        "Order Date",
        "Status",
        "Product SKU",
        "Item ID",
        "Item Name",
        "Item Quantity",
        "Selling Price",
        "Synthetic_Data_Flag",
    }

    workbook = load_workbook(
        SALES_FILE, read_only=True, data_only=True
    )

    try:
        if "Sales Data" not in workbook.sheetnames:
            raise ValueError("Worksheet 'Sales Data' not found.")

        sheet = workbook["Sales Data"]
        header = [
            cell.value
            for cell in next(
                sheet.iter_rows(min_row=1, max_row=1)
            )
        ]

        missing = required_columns - set(header)
        if missing:
            raise ValueError(
                f"Sales workbook missing columns: {sorted(missing)}"
            )

        positions = {name: header.index(name) for name in required_columns}
        total = 0
        errors = []

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if not any(value is not None for value in row):
                continue

            try:
                quantity = row[positions["Item Quantity"]]
                product_sku = row[positions["Product SKU"]]
                item_id = row[positions["Item ID"]]
                status = row[positions["Status"]]

                if quantity is None or int(quantity) < 0:
                    raise ValueError("Quantity is missing or negative.")

                if not product_sku and not item_id:
                    raise ValueError("Both product identifiers are blank.")

                if not status:
                    raise ValueError("Order status is blank.")

                total += 1

            except (ValueError, TypeError, IndexError) as exc:
                errors.append(f"Row {row_number}: {exc}")

        print(f"Sales rows checked: {total + len(errors)}")
        print(f"Sales valid rows: {total}")
        print(f"Sales errors: {len(errors)}")

        for error in errors:
            print(error)

        return total, len(errors)

    finally:
        workbook.close()


def main():
    if not INVENTORY_FILE.exists():
        raise FileNotFoundError(INVENTORY_FILE)

    if not SALES_FILE.exists():
        raise FileNotFoundError(SALES_FILE)

    print("VALIDATING INVENTORY REPORT")
    inventory_valid, inventory_errors = validate_inventory()

    print("\nVALIDATING SALES REPORT")
    sales_valid, sales_errors = validate_sales()

    print("\nVALIDATION SUMMARY")
    print(f"Inventory valid: {inventory_valid}")
    print(f"Sales valid: {sales_valid}")
    print(f"Total validation errors: {inventory_errors + sales_errors}")
    print("No database changes were made.")


if __name__ == "__main__":
    main()
